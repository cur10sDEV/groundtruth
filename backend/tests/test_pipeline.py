import asyncio
import hashlib
import io
import uuid
from contextlib import contextmanager
from unittest.mock import AsyncMock

import pytest
from sqlalchemy import select

import app.db as db_module
from app.core.errors import IngestionError
from app.core.metrics import INGESTION_FAILED, INGESTION_PROCESSED
from app.db import get_session, init_db
from app.models.chunk import Chunk
from app.models.document import Document, DocumentStatus
from app.rag.chunkers.base import ChunkData
from app.rag.embed.embeddings import SparseVector
from app.rag.ingestion import pipeline
from app.rag.ingestion.pipeline import (
    cleanup_stale,
    ingest_document,
    ingest_versioned,
)

DOC_ID = "11111111-1111-1111-1111-111111111111"
USER_ID = "22222222-2222-2222-2222-222222222222"
ORG_ID = "33333333-3333-3333-3333-333333333333"
S3_KEY = "documents/org/user/doc/sample.txt"

REQUIRED_PAYLOAD_KEYS = {
    "doc_id",
    "user_id",
    "org_id",
    "chunk_text_hash",
    "version",
    "page_number",
    "chunk_index",
    "word_count",
    "char_count",
    "type",
}


def _doc_version_filter(version: int) -> dict:
    return {
        "must": [
            {"key": "doc_id", "match": {"value": DOC_ID}},
            {"key": "version", "match": {"value": version}},
        ]
    }


@pytest.mark.integration
async def test_ingest_document_returns_chunk_count():
    # requires running MinIO, Postgres, Qdrant and a pre-seeded Document row
    count = await ingest_document("doc-id-1", "documents/org/user/doc-id-1/sample.txt")
    assert count >= 1


@pytest.fixture
def _sqlite_database_url(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite+aiosqlite:///{tmp_path}/test.db")
    monkeypatch.setattr(db_module, "_engine", None)
    monkeypatch.setattr(db_module, "_sessionmaker", None)


@pytest.fixture
async def db(_sqlite_database_url):
    await init_db()
    yield
    if db_module._engine is not None:
        await db_module._engine.dispose()


@pytest.fixture
def fake_s3(monkeypatch):
    objects: dict[str, bytes] = {}

    def fake_get_object(s3_key: str) -> bytes:
        return objects[s3_key]

    monkeypatch.setattr(pipeline, "get_object", fake_get_object)
    return objects


@pytest.fixture
def fake_qdrant(monkeypatch):
    calls = {"upserts": [], "deletes": [], "points": {}}

    def fake_upsert_points_batch(points):
        calls["upserts"].extend(points)
        for p in points:
            calls["points"][p["id"]] = p["payload"]

    def fake_delete_points(point_ids, payload_filter=None):
        calls["deletes"].append((point_ids, payload_filter))
        if point_ids:
            for point_id in point_ids:
                calls["points"].pop(point_id, None)
        elif payload_filter:
            must = {m["key"]: m["match"]["value"] for m in payload_filter.get("must", [])}
            for point_id, payload in list(calls["points"].items()):
                if all(payload.get(k) == v for k, v in must.items()):
                    del calls["points"][point_id]

    monkeypatch.setattr(pipeline, "upsert_points_batch", fake_upsert_points_batch)
    monkeypatch.setattr(pipeline, "delete_points", fake_delete_points)
    return calls


@pytest.fixture
def fake_embeddings(monkeypatch):
    calls = {"dense": [], "sparse": []}

    async def fake_dense_embed(texts):
        calls["dense"].append(list(texts))
        return [[float(len(t)), 1.0] for t in texts]

    def fake_sparse_embed(texts):
        calls["sparse"].append(list(texts))
        return [SparseVector(indices=[1], values=[1.0]) for _ in texts]

    monkeypatch.setattr(pipeline, "dense_embed", fake_dense_embed)
    monkeypatch.setattr(pipeline, "sparse_embed", fake_sparse_embed)
    return calls


@pytest.fixture
def fake_cache(monkeypatch):
    calls: list[tuple[str, str]] = []

    async def fake_invalidate_for_doc(org_id: str, doc_id: str) -> int:
        calls.append((org_id, doc_id))
        return 0

    monkeypatch.setattr(pipeline, "invalidate_for_doc", fake_invalidate_for_doc)
    return calls


async def _seed_doc(
    filename: str = "sample.txt",
    content_hash: str = "stale-hash",
    current_version: int = 1,
) -> None:
    async with get_session() as session:
        session.add(
            Document(
                id=DOC_ID,
                user_id=USER_ID,
                org_id=ORG_ID,
                original_filename=filename,
                status=DocumentStatus.PENDING,
                content_hash=content_hash,
                current_version=current_version,
            )
        )
        await session.commit()


async def _seed_chunk(version: int, text: str = "some text") -> str:
    async with get_session() as session:
        chunk = Chunk(
            doc_id=DOC_ID,
            user_id=USER_ID,
            org_id=ORG_ID,
            chunk_text=text,
            page_number=0,
            start_offset=0,
            end_offset=len(text),
            version=version,
        )
        session.add(chunk)
        await session.commit()
        return chunk.id


async def _get_doc() -> Document:
    async with get_session() as session:
        return await session.get(Document, DOC_ID)


async def _chunks_for(doc_id: str = DOC_ID) -> list[Chunk]:
    async with get_session() as session:
        rows = (await session.execute(select(Chunk).where(Chunk.doc_id == doc_id))).scalars().all()
        return list(rows)


def _build_pdf() -> bytes:
    import fitz

    doc = fitz.open()
    for text in ("Alpha beta gamma delta.", "Epsilon zeta eta theta."):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _build_docx() -> bytes:
    from docx import Document as DocxDocument

    buf = io.BytesIO()
    d = DocxDocument()
    d.add_paragraph("Alpha beta gamma delta.")
    d.add_paragraph("Epsilon zeta eta theta.")
    d.save(buf)
    return buf.getvalue()


async def test_ingest_document_persists_chunks_and_payload(
    db, fake_s3, fake_qdrant, fake_embeddings
):
    body = b"Alpha beta gamma delta epsilon zeta eta theta."
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash="different")

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count >= 1
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.content_hash == hashlib.sha256(body).hexdigest()

    rows = await _chunks_for()
    assert len(rows) == count
    rows_by_id = {}
    for r in rows:
        assert r.doc_id == DOC_ID
        assert r.user_id == USER_ID
        assert r.org_id == ORG_ID
        assert r.version == 1
        assert r.chunk_text
        rows_by_id[r.id] = r

    assert len(fake_qdrant["upserts"]) == count
    assert {p["id"] for p in fake_qdrant["upserts"]} == set(rows_by_id)
    for point in fake_qdrant["upserts"]:
        payload = point["payload"]
        assert set(payload) >= REQUIRED_PAYLOAD_KEYS
        row = rows_by_id[point["id"]]
        assert payload["doc_id"] == DOC_ID
        assert payload["user_id"] == USER_ID
        assert payload["org_id"] == ORG_ID
        assert payload["version"] == 1
        assert payload["type"] == "txt"
        assert payload["chunk_text_hash"] == hashlib.sha256(row.chunk_text.encode()).hexdigest()
        assert payload["word_count"] == len(row.chunk_text.split())
        assert payload["char_count"] == len(row.chunk_text)
        assert len(point["dense"]) == 2
        assert point["sparse_indices"] == [1]
        assert point["sparse_values"] == [1.0]
    assert fake_embeddings["dense"]


async def test_ingest_document_dedup_returns_zero(db, fake_s3, fake_qdrant, fake_embeddings):
    body = b"same bytes as before"
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash=hashlib.sha256(body).hexdigest())

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count == 0
    assert (await _get_doc()).status == DocumentStatus.EMBEDDED
    assert fake_qdrant["upserts"] == []
    assert fake_embeddings["dense"] == []
    assert await _chunks_for() == []


async def test_ingest_document_missing_doc_raises(db):
    with pytest.raises(IngestionError):
        await ingest_document("missing-doc-id", S3_KEY)


async def test_ingest_document_wraps_parser_errors(db, fake_s3, fake_qdrant, fake_embeddings):
    fake_s3[S3_KEY] = b""
    await _seed_doc(filename="broken.pdf", content_hash="different")

    with pytest.raises(IngestionError) as excinfo:
        await ingest_document(DOC_ID, S3_KEY)

    assert "parse failed" in excinfo.value.detail
    assert excinfo.value.__cause__ is not None
    assert not isinstance(excinfo.value.__cause__, IngestionError)
    assert (await _get_doc()).status == DocumentStatus.FAILED


async def test_ingest_document_skips_marker_only_chunks(db, fake_s3, fake_qdrant, fake_embeddings):
    fake_s3[S3_KEY] = b"# Intro\n\nBody text goes here.\n\n\n\n"
    await _seed_doc(content_hash="different")

    count = await ingest_document(DOC_ID, S3_KEY)

    rows = await _chunks_for()
    assert count == len(rows) == 2
    assert {r.chunk_text for r in rows} == {"Intro", "Body text goes here."}
    assert len(fake_qdrant["upserts"]) == 2


async def test_ingest_document_all_chunks_empty_returns_zero(
    db, fake_s3, fake_qdrant, fake_embeddings
):
    fake_s3[S3_KEY] = b"   \n\n \t \n"
    await _seed_doc(content_hash="different")

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count == 0
    assert await _chunks_for() == []
    assert fake_qdrant["upserts"] == []
    assert fake_embeddings["dense"] == []
    assert (await _get_doc()).status == DocumentStatus.EMBEDDED


async def test_ingest_document_marks_failed_on_embedding_error(
    db, fake_s3, fake_qdrant, monkeypatch
):
    fake_s3[S3_KEY] = b"Alpha beta gamma delta."
    await _seed_doc(content_hash="different")

    async def failing_dense_embed(texts):
        raise RuntimeError("embed down")

    monkeypatch.setattr(pipeline, "dense_embed", failing_dense_embed)
    with pytest.raises(RuntimeError):
        await ingest_document(DOC_ID, S3_KEY)

    assert (await _get_doc()).status == DocumentStatus.FAILED
    assert await _chunks_for() == []


async def test_ingest_cancelled_marks_failed_not_processing(db, fake_s3, fake_qdrant, monkeypatch):
    fake_s3[S3_KEY] = b"Alpha beta gamma delta."
    await _seed_doc(content_hash="different")

    async def cancelled_dense_embed(texts):
        raise asyncio.CancelledError

    monkeypatch.setattr(pipeline, "dense_embed", cancelled_dense_embed)

    with pytest.raises(asyncio.CancelledError):
        await ingest_document(DOC_ID, S3_KEY)

    doc = await _get_doc()
    assert doc.status == DocumentStatus.FAILED
    assert doc.pending_version is None
    # only the clean-slate reset ran; the cancel triggered no orphan cleanup
    assert fake_qdrant["deletes"] == [([], _doc_version_filter(1))]


async def test_ingest_versioned_sets_pending_version_during_attempt(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache, monkeypatch
):
    fake_s3[S3_KEY] = b"Version two content with fresh words."
    await _seed_doc(current_version=1)

    observed = []

    async def spying_dense_embed(texts):
        doc = await _get_doc()
        observed.append((doc.status, doc.pending_version))
        return [[float(len(t)), 1.0] for t in texts]

    monkeypatch.setattr(pipeline, "dense_embed", spying_dense_embed)

    await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert observed
    assert {(status, pending) for status, pending in observed} == {(DocumentStatus.PROCESSING, 2)}
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.pending_version is None


async def test_ingest_failure_clears_pending_version(db, fake_s3, fake_qdrant, monkeypatch):
    fake_s3[S3_KEY] = b"Alpha beta gamma delta."
    await _seed_doc(content_hash="different", current_version=1)

    async def failing_dense_embed(texts):
        raise RuntimeError("embed down")

    monkeypatch.setattr(pipeline, "dense_embed", failing_dense_embed)

    with pytest.raises(RuntimeError):
        await ingest_document(DOC_ID, S3_KEY)

    doc = await _get_doc()
    assert doc.status == DocumentStatus.FAILED
    assert doc.pending_version is None


async def test_ingest_failure_deletes_partial_qdrant_points(
    db, fake_s3, fake_qdrant, fake_embeddings, monkeypatch
):
    body = "\n\n".join(f"Paragraph {i} with a few words." for i in range(40)).encode()
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash="different")

    upsert_calls = []

    def failing_upsert_points_batch(points):
        upsert_calls.append(len(points))
        if len(upsert_calls) >= 2:
            raise RuntimeError("qdrant down mid-batch")

    monkeypatch.setattr(pipeline, "upsert_points_batch", failing_upsert_points_batch)

    with pytest.raises(RuntimeError):
        await ingest_document(DOC_ID, S3_KEY)

    assert upsert_calls == [32, 8]
    assert (await _get_doc()).status == DocumentStatus.FAILED
    assert await _chunks_for() == []
    # clean-slate reset at attempt start, then orphan cleanup after the failure
    assert len(fake_qdrant["deletes"]) == 2
    for point_ids, payload_filter in fake_qdrant["deletes"]:
        assert point_ids == []
        assert payload_filter == _doc_version_filter(1)


async def test_ingest_pdf_roundtrip_per_page(db, fake_s3, fake_qdrant, fake_embeddings):
    fake_s3[S3_KEY] = _build_pdf()
    await _seed_doc(filename="doc.pdf", content_hash="different")

    count = await ingest_document(DOC_ID, S3_KEY)

    rows = await _chunks_for()
    assert count == len(rows) == 2
    assert sorted(r.page_number for r in rows) == [0, 1]
    texts_by_page = {r.page_number: r.chunk_text for r in rows}
    assert texts_by_page[0] == "Alpha beta gamma delta."
    assert texts_by_page[1] == "Epsilon zeta eta theta."
    for point in fake_qdrant["upserts"]:
        assert point["payload"]["type"] == "pdf"
        assert point["payload"]["page_number"] in (0, 1)


async def test_ingest_docx_roundtrip_single_page(db, fake_s3, fake_qdrant, fake_embeddings):
    fake_s3[S3_KEY] = _build_docx()
    await _seed_doc(filename="doc.docx", content_hash="different")

    count = await ingest_document(DOC_ID, S3_KEY)

    rows = await _chunks_for()
    assert count == len(rows) == 1
    assert rows[0].page_number == 0
    assert "Alpha beta gamma delta." in rows[0].chunk_text
    assert fake_qdrant["upserts"][0]["payload"]["type"] == "docx"


async def test_ingest_versioned_writes_new_version_then_flips(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache
):
    fake_s3[S3_KEY] = b"Version two content with fresh words."
    await _seed_doc(current_version=1)

    count = await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert count >= 1
    rows = await _chunks_for()
    assert len(rows) == count
    assert {r.version for r in rows} == {2}
    doc = await _get_doc()
    assert doc.current_version == 2
    assert doc.status == DocumentStatus.EMBEDDED
    assert fake_cache == [(ORG_ID, DOC_ID)]
    for point in fake_qdrant["upserts"]:
        assert point["payload"]["version"] == 2


async def test_ingest_versioned_reingests_unchanged_content_at_new_version(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache
):
    body = b"Identical content across versions."
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash=hashlib.sha256(body).hexdigest(), current_version=1)

    count = await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert count >= 1
    assert {r.version for r in await _chunks_for()} == {2}
    assert (await _get_doc()).current_version == 2
    assert fake_cache == [(ORG_ID, DOC_ID)]


async def test_ingest_versioned_failure_keeps_current_version(
    db, fake_s3, fake_qdrant, fake_cache, monkeypatch
):
    fake_s3[S3_KEY] = b"Version two content."
    await _seed_doc(current_version=1)

    async def failing_dense_embed(texts):
        raise RuntimeError("embed down")

    monkeypatch.setattr(pipeline, "dense_embed", failing_dense_embed)
    with pytest.raises(RuntimeError):
        await ingest_versioned(DOC_ID, S3_KEY, 2)

    doc = await _get_doc()
    assert doc.current_version == 1
    assert doc.status == DocumentStatus.FAILED
    assert await _chunks_for() == []
    assert fake_cache == []


async def test_versioned_flip_keeps_v1_serving_until_flip(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache, monkeypatch
):
    from app.rag.retrieval.orchestrator import resolve_text_for_chunk_ids

    fake_s3[S3_KEY] = b"Version two content with fresh words."
    # PENDING is the real state of a versioned re-ingest: the upload route flips
    # the doc EMBEDDED → PENDING when publishing, and only the guarded finalize
    # (which also rolls current_version) may set it back to EMBEDDED.
    await _seed_doc(current_version=1)
    v1_id = await _seed_chunk(version=1, text="v1 body")

    # pre-flip: v1 is the current version and serves
    assert await resolve_text_for_chunk_ids([v1_id], ORG_ID) == [{"id": v1_id, "text": "v1 body"}]

    served_mid_ingest = []

    async def spying_dense_embed(texts):
        # mid-ingest: v2 chunks are not committed and the flip is not applied,
        # so retrieval-side resolution must still serve v1
        served_mid_ingest.append(await resolve_text_for_chunk_ids([v1_id], ORG_ID))
        return [[float(len(t)), 1.0] for t in texts]

    monkeypatch.setattr(pipeline, "dense_embed", spying_dense_embed)

    await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert served_mid_ingest == [[{"id": v1_id, "text": "v1 body"}]]

    # post-flip: v1 is stale (mixed window until cleanup) and must not serve
    assert await resolve_text_for_chunk_ids([v1_id], ORG_ID) == []
    v2_ids = [r.id for r in await _chunks_for() if r.version == 2]
    assert v2_ids
    resolved_v2 = await resolve_text_for_chunk_ids(v2_ids, ORG_ID)
    assert {r["id"] for r in resolved_v2} == set(v2_ids)


async def test_cleanup_stale_deletes_old_versions(db, fake_qdrant):
    await _seed_doc(current_version=2)
    stale_id = await _seed_chunk(version=1, text="old text")
    kept_id = await _seed_chunk(version=2, text="new text")

    removed = await cleanup_stale(DOC_ID)

    assert removed == 1
    assert {r.id for r in await _chunks_for()} == {kept_id}
    assert fake_qdrant["deletes"] == [
        ([stale_id], {"must": [{"key": "doc_id", "match": {"value": DOC_ID}}]})
    ]


async def test_cleanup_stale_missing_doc_returns_zero(db, fake_qdrant):
    assert await cleanup_stale("missing-doc-id") == 0
    assert fake_qdrant["deletes"] == []


# --- Prometheus metrics (Task 6.3 wiring) ---


async def test_ingest_success_increments_processed_total(db, fake_s3, fake_qdrant, fake_embeddings):
    fake_s3[S3_KEY] = b"Alpha beta gamma delta epsilon."
    await _seed_doc(content_hash="different")

    before = INGESTION_PROCESSED._value.get()

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count >= 1
    assert INGESTION_PROCESSED._value.get() == before + 1


async def test_ingest_dedup_completing_embedded_counts_as_processed(
    db, fake_s3, fake_qdrant, fake_embeddings
):
    body = b"same bytes as before"
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash=hashlib.sha256(body).hexdigest())

    before = INGESTION_PROCESSED._value.get()

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count == 0
    assert INGESTION_PROCESSED._value.get() == before + 1


async def test_ingest_failure_increments_failed_total(db, fake_s3, fake_qdrant, monkeypatch):
    fake_s3[S3_KEY] = b"Alpha beta gamma delta."
    await _seed_doc(content_hash="different")

    async def failing_dense_embed(texts):
        raise RuntimeError("embed down")

    monkeypatch.setattr(pipeline, "dense_embed", failing_dense_embed)

    before = INGESTION_FAILED._value.get()

    with pytest.raises(RuntimeError):
        await ingest_document(DOC_ID, S3_KEY)

    assert INGESTION_FAILED._value.get() == before + 1


# --- Deterministic chunk ids / idempotent re-ingest ---


@contextmanager
def _crash_on_finalize():
    """Simulate a hard crash after the chunk-rows commit, before the guarded finalize."""
    real_finalize = pipeline._finalize

    async def crashing_finalize(doc_id, content_hash, new_version=None):
        raise KeyboardInterrupt

    pipeline._finalize = crashing_finalize
    try:
        yield
    finally:
        pipeline._finalize = real_finalize


def test_chunk_id_derivation_is_deterministic():
    first = pipeline._derive_chunk_id(DOC_ID, 2, 0)

    assert first == pipeline._derive_chunk_id(DOC_ID, 2, 0)
    assert len(first) == 36
    uuid.UUID(first)
    assert pipeline._derive_chunk_id(DOC_ID, 3, 0) != first
    assert pipeline._derive_chunk_id(DOC_ID, 2, 1) != first
    assert pipeline._derive_chunk_id("44444444-4444-4444-4444-444444444444", 2, 0) != first


async def test_versioned_crash_window_redelivery_converges(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache
):
    fake_s3[S3_KEY] = b"Version two content with fresh words."
    await _seed_doc(current_version=1)

    with _crash_on_finalize(), pytest.raises(KeyboardInterrupt):
        await ingest_versioned(DOC_ID, S3_KEY, 2)

    # crash window state: v2 chunks/points committed but the flip never happened
    doc = await _get_doc()
    assert doc.current_version == 1
    crashed_rows = await _chunks_for()
    assert len(crashed_rows) >= 1
    assert {r.version for r in crashed_rows} == {2}
    assert {p["id"] for p in fake_qdrant["upserts"]} == {r.id for r in crashed_rows}
    assert len(fake_qdrant["points"]) == len(crashed_rows)

    # redelivery of the same message
    count = await ingest_versioned(DOC_ID, S3_KEY, 2)

    doc = await _get_doc()
    assert doc.current_version == 2
    assert doc.status == DocumentStatus.EMBEDDED
    rows = await _chunks_for()
    assert len(rows) == count
    expected_ids = {pipeline._derive_chunk_id(DOC_ID, 2, i) for i in range(count)}
    assert {r.id for r in rows} == expected_ids
    # exactly one set of points survives: the derived ids, no duplicates
    assert set(fake_qdrant["points"]) == expected_ids
    assert len(fake_qdrant["upserts"]) == 2 * count
    assert len({p["id"] for p in fake_qdrant["upserts"]}) == count
    assert fake_qdrant["deletes"] == [([], _doc_version_filter(2)), ([], _doc_version_filter(2))]
    assert fake_cache == [(ORG_ID, DOC_ID)]


async def test_versioned_retry_with_fewer_chunks_leaves_no_leftovers(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache, monkeypatch
):
    body = "\n\n".join(f"Paragraph {i} carries fresh words." for i in range(3)).encode()
    fake_s3[S3_KEY] = body
    await _seed_doc(current_version=1)

    with _crash_on_finalize(), pytest.raises(KeyboardInterrupt):
        await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert len(await _chunks_for()) == 3
    assert len(fake_qdrant["points"]) == 3

    class TwoChunkChunker:
        def chunk(self, text, page_number=0):
            return [
                ChunkData(text="Retry chunk one.", page_number=page_number, chunk_index=0),
                ChunkData(text="Retry chunk two.", page_number=page_number, chunk_index=1),
            ]

    monkeypatch.setattr(pipeline, "get_chunker", lambda _ext: TwoChunkChunker())

    count = await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert count == 2
    expected_ids = {
        pipeline._derive_chunk_id(DOC_ID, 2, 0),
        pipeline._derive_chunk_id(DOC_ID, 2, 1),
    }
    assert {r.id for r in await _chunks_for()} == expected_ids
    # the leftover third point/row from attempt 1 did not survive the retry
    assert set(fake_qdrant["points"]) == expected_ids
    assert pipeline._derive_chunk_id(DOC_ID, 2, 2) not in fake_qdrant["points"]
    doc = await _get_doc()
    assert doc.current_version == 2
    assert doc.status == DocumentStatus.EMBEDDED


async def test_dedup_short_circuit_does_not_wipe_existing_chunks(
    db, fake_s3, fake_qdrant, fake_embeddings
):
    body = b"same bytes as before"
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash=hashlib.sha256(body).hexdigest())
    kept_id = await _seed_chunk(version=1, text="existing chunk body")

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count == 0
    assert {r.id for r in await _chunks_for()} == {kept_id}
    assert fake_qdrant["upserts"] == []
    assert fake_qdrant["deletes"] == []


async def test_plain_path_crash_retry_converges(
    db, fake_s3, fake_qdrant, fake_embeddings, monkeypatch
):
    body = b"Alpha beta gamma delta epsilon zeta eta theta."
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash="different")

    recorded_upsert = pipeline.upsert_points_batch

    def crash_after_upsert(points):
        recorded_upsert(points)
        raise KeyboardInterrupt

    monkeypatch.setattr(pipeline, "upsert_points_batch", crash_after_upsert)
    with pytest.raises(KeyboardInterrupt):
        await ingest_document(DOC_ID, S3_KEY)

    # hard-crash residue: orphan points in qdrant, no PG rows, doc stuck in PROCESSING
    assert (await _get_doc()).status == DocumentStatus.PROCESSING
    assert await _chunks_for() == []
    orphan_ids = set(fake_qdrant["points"])
    assert orphan_ids
    assert fake_qdrant["deletes"] == [([], _doc_version_filter(1))]

    monkeypatch.setattr(pipeline, "upsert_points_batch", recorded_upsert)

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count >= 1
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.pending_version is None
    assert doc.content_hash == hashlib.sha256(body).hexdigest()
    rows = await _chunks_for()
    assert len(rows) == count
    expected_ids = {pipeline._derive_chunk_id(DOC_ID, 1, i) for i in range(count)}
    assert {r.id for r in rows} == expected_ids
    # the orphans were overwritten by the same deterministic ids: one set remains
    assert set(fake_qdrant["points"]) == expected_ids
    assert orphan_ids == expected_ids
    assert fake_qdrant["deletes"] == [([], _doc_version_filter(1)), ([], _doc_version_filter(1))]


# --- Guarded status transitions: deletion must be un-raceable ---


async def test_delete_during_ingest_cannot_be_flipped(db):
    # simulate: a DELETING row while _ingest runs
    await _seed_doc()
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.DELETING
        await session.commit()

    # claim must fail: DELETING row cannot be re-claimed
    with pytest.raises(IngestionError):
        await pipeline._ingest(DOC_ID, S3_KEY)


async def test_finalize_loses_race_to_deleting(db):
    # doc PROCESSING, chunks committed, then deletion flips status mid-run
    await _seed_doc()
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.PROCESSING
        doc.pending_version = 1
        await session.commit()
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.DELETING
        await session.commit()

    with pytest.raises(IngestionError):
        await pipeline._finalize(DOC_ID, "hash123")

    doc = await _get_doc()
    assert doc.status == DocumentStatus.DELETING  # not stomped


async def test_mark_failed_does_not_stomp_deleting(db):
    await _seed_doc()
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.DELETING
        await session.commit()

    await pipeline._mark_failed(DOC_ID, reason="boom")

    doc = await _get_doc()
    assert doc.status == DocumentStatus.DELETING
    assert doc.failure_reason is None


async def test_mark_failed_persists_reason_when_landed(db):
    await _seed_doc()
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.PROCESSING
        await session.commit()

    landed = await pipeline._mark_failed(DOC_ID, reason="boom")

    assert landed is True
    doc = await _get_doc()
    assert doc.status == DocumentStatus.FAILED
    assert doc.failure_reason == "boom"
    assert doc.pending_version is None


async def test_mark_failed_lost_race_keeps_failed_metric_flat(db):
    await _seed_doc()
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.DELETING
        await session.commit()

    before = INGESTION_FAILED._value.get()

    landed = await pipeline._mark_failed(DOC_ID, reason="boom")

    assert landed is False
    assert INGESTION_FAILED._value.get() == before


async def test_finalize_clears_stale_failure_reason_on_recovery(
    db, fake_s3, fake_qdrant, fake_embeddings
):
    fake_s3[S3_KEY] = b"Alpha beta gamma delta epsilon."
    await _seed_doc(content_hash="different")
    async with get_session() as session:
        doc = await session.get(Document, DOC_ID)
        doc.status = DocumentStatus.FAILED
        doc.failure_reason = "previous attempt broke"
        await session.commit()

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count >= 1
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.failure_reason is None


# --- Cross-document duplicate detection (user-scoped) ---


OTHER_DOC_ID = "44444444-4444-4444-4444-444444444444"
OTHER_USER_ID = "55555555-5555-5555-5555-555555555555"


async def _seed_other_doc(
    doc_id: str = OTHER_DOC_ID,
    user_id: str = USER_ID,
    status: DocumentStatus = DocumentStatus.EMBEDDED,
    content_hash: str = "other-hash",
) -> None:
    async with get_session() as session:
        session.add(
            Document(
                id=doc_id,
                user_id=user_id,
                org_id=ORG_ID,
                original_filename="twin.txt",
                status=status,
                content_hash=content_hash,
                current_version=1,
            )
        )
        await session.commit()


async def test_cross_doc_duplicate_marks_failed_with_reason(
    db, fake_s3, fake_qdrant, fake_embeddings, monkeypatch
):
    body = b"Alpha beta gamma delta epsilon zeta eta theta."
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash="different")
    await _seed_other_doc(content_hash=hashlib.sha256(body).hexdigest())

    deleted: list[str] = []
    monkeypatch.setattr(pipeline, "delete_object", lambda k: deleted.append(k))
    reset = AsyncMock()
    monkeypatch.setattr(pipeline, "_reset_version", reset)
    failed_before = INGESTION_FAILED._value.get()
    processed_before = INGESTION_PROCESSED._value.get()

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count == 0
    doc = await _get_doc()
    assert doc.status == DocumentStatus.FAILED
    assert doc.pending_version is None
    assert doc.failure_reason == f"duplicate of {OTHER_DOC_ID}"
    assert deleted == [S3_KEY]
    reset.assert_not_called()
    assert fake_qdrant["upserts"] == []
    assert fake_embeddings["dense"] == []
    assert await _chunks_for() == []
    # a duplicate is an expected outcome, not a system failure: no metric moves
    assert INGESTION_FAILED._value.get() == failed_before
    assert INGESTION_PROCESSED._value.get() == processed_before


async def test_duplicate_check_is_user_scoped(
    db, fake_s3, fake_qdrant, fake_embeddings, monkeypatch
):
    body = b"Alpha beta gamma delta epsilon zeta eta theta."
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash="different")
    # same org, DIFFERENT user, same bytes: not this user's duplicate
    await _seed_other_doc(user_id=OTHER_USER_ID, content_hash=hashlib.sha256(body).hexdigest())

    deleted: list[str] = []
    monkeypatch.setattr(pipeline, "delete_object", lambda k: deleted.append(k))

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count >= 1
    assert deleted == []
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.failure_reason is None


async def test_duplicate_check_excludes_self(
    db, fake_s3, fake_qdrant, fake_embeddings, fake_cache, monkeypatch
):
    body = b"Identical content across versions."
    fake_s3[S3_KEY] = body
    # doc already served these exact bytes at v1; republishing them as v2
    # must not flag the doc as a duplicate of itself
    await _seed_doc(content_hash=hashlib.sha256(body).hexdigest(), current_version=1)

    deleted: list[str] = []
    monkeypatch.setattr(pipeline, "delete_object", lambda k: deleted.append(k))

    count = await ingest_versioned(DOC_ID, S3_KEY, 2)

    assert count >= 1
    assert deleted == []
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.current_version == 2
    assert doc.failure_reason is None
    assert {r.version for r in await _chunks_for()} == {2}


async def test_duplicate_ignores_non_embedded(
    db, fake_s3, fake_qdrant, fake_embeddings, monkeypatch
):
    body = b"Alpha beta gamma delta epsilon zeta eta theta."
    fake_s3[S3_KEY] = body
    await _seed_doc(content_hash="different")
    # same user, same bytes, but never finished ingesting: not a duplicate source
    await _seed_other_doc(
        status=DocumentStatus.PENDING, content_hash=hashlib.sha256(body).hexdigest()
    )

    deleted: list[str] = []
    monkeypatch.setattr(pipeline, "delete_object", lambda k: deleted.append(k))

    count = await ingest_document(DOC_ID, S3_KEY)

    assert count >= 1
    assert deleted == []
    doc = await _get_doc()
    assert doc.status == DocumentStatus.EMBEDDED
    assert doc.failure_reason is None
