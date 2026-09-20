from dataclasses import dataclass

from app.core.errors import IngestionError


@dataclass
class ParsedPage:
    text: str
    page_number: int


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    full_text: str


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise IngestionError(detail="PyMuPDF not installed") from exc
    pages: list[ParsedPage] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            pages.append(ParsedPage(text=page.get_text(), page_number=i))
    return ParsedDocument(pages=pages, full_text="\n".join(p.text for p in pages))


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise IngestionError(detail="python-docx not installed") from exc
    import io

    doc = DocxDocument(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    return ParsedDocument(pages=[ParsedPage(text=text, page_number=0)], full_text=text)


def parse_markdown(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    return ParsedDocument(pages=[ParsedPage(text=text, page_number=0)], full_text=text)


def parse_txt(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
    return ParsedDocument(pages=[ParsedPage(text=text, page_number=0)], full_text=text)


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return parse_pdf(data)
    if ext == "docx":
        return parse_docx(data)
    if ext == "md":
        return parse_markdown(data)
    if ext == "txt":
        return parse_txt(data)
    raise IngestionError(detail=f"unsupported file type: {ext or 'unknown'}")
